import openpyxl
from datetime import datetime, timedelta, date
import calendar
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.utils import timezone
from django.db.models import Q, Count, Case, When, IntegerField
from django.db.models.functions import TruncMonth
from django.http import HttpResponse
from django.urls import reverse

from .models import (
    Student,
    Asrama,
    Kelas,
    Absensi,
    Guru,
    PresensiGuru,
    TahunAjaran,
    MataPelajaran,
    Penilaian,
)


# =========================================================
# DASHBOARD
# =========================================================

def dashboard(request):
    total_siswa = Student.objects.count()
    total_kelas = Kelas.objects.count()
    total_mata_pelajaran = MataPelajaran.objects.count()
    tahun_ajaran = TahunAjaran.objects.filter(aktif=True).first()

    context = {
        'total_siswa': total_siswa,
        'total_guru': 0,
        'total_kelas': total_kelas,
        'total_mata_pelajaran': total_mata_pelajaran,
        'tahun_ajaran': tahun_ajaran,
    }

    return render(
        request,
        'students/dashboard.html',
        context
    )


# =========================================================
# DATA SISWA
# =========================================================

def student_list(request):

    search = request.GET.get(
        'search',
        ''
    ).strip()

    tahun_ajaran_id = request.GET.get(
        'tahun_ajaran',
        ''
    ).strip()

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    asrama_id = request.GET.get(
        'asrama',
        ''
    ).strip()

    status = request.GET.get(
        'status',
        ''
    ).strip()

    students = (
        Student.objects
        .select_related(
            'kelas',
            'asrama_master',
            'tahun_ajaran'
        )
        .all()
        .order_by('nama')
    )

    # =========================================
    # SEARCH NAMA / NIM
    # =========================================

    if search:

        students = students.filter(
            Q(nama__icontains=search) |
            Q(nim__icontains=search)
        )

    # =========================================
    # FILTER TAHUN AJARAN
    # =========================================

    if tahun_ajaran_id:

        students = students.filter(
            tahun_ajaran_id=tahun_ajaran_id
        )

    # =========================================
    # FILTER KELAS
    # =========================================

    if kelas_id:

        students = students.filter(
            kelas_id=kelas_id
        )

    # =========================================
    # FILTER ASRAMA
    # =========================================

    if asrama_id:

        students = students.filter(
            asrama_master_id=asrama_id
        )

    # =========================================
    # FILTER STATUS
    # =========================================

    if status == 'aktif':

        students = students.filter(
            status=True
        )

    elif status == 'nonaktif':

        students = students.filter(
            status=False
        )

    # =========================================
    # DATA DROPDOWN
    # =========================================

    tahun_ajarans = (
        TahunAjaran.objects
        .all()
        .order_by('-nama')
    )

    kelass = (
        Kelas.objects
        .all()
        .order_by('nama')
    )

    asramas = (
        Asrama.objects
        .all()
        .order_by('nama')
    )

    # =========================================
    # CONTEXT
    # =========================================

    context = {

        'students':
            students,

        'search':
            search,

        'tahun_ajarans':
            tahun_ajarans,

        'kelass':
            kelass,

        'asramas':
            asramas,

        'tahun_ajaran_id':
            tahun_ajaran_id,

        'kelas_id':
            kelas_id,

        'asrama_id':
            asrama_id,

        'status':
            status,
    }

    return render(
        request,
        'students/student_list.html',
        context
    )

def student_detail(request, id):
    student = get_object_or_404(
        Student.objects.select_related(
            'kelas',
            'asrama_master',
            'tahun_ajaran'
        ),
        id=id
    )

    return render(
        request,
        'students/student_detail.html',
        {
            'student': student,
        }
    )

# =========================================================
# TAMBAH SISWA
# =========================================================

def student_create(request):

    asramas = (
        Asrama.objects
        .all()
        .order_by('nama')
    )

    kelass = (
        Kelas.objects
        .all()
        .order_by('nama')
    )

    tahun_ajarans = (
        TahunAjaran.objects
        .all()
        .order_by('-nama')
    )

    if request.method == 'POST':

        nama = request.POST.get(
            'nama',
            ''
        ).strip().title()

        nama_ayah = request.POST.get(
            'nama_ayah',
            ''
        ).strip().title()

        nim = request.POST.get(
            'nim',
            ''
        ).strip()

        jk = request.POST.get(
            'jk',
            ''
        ).strip()

        tempat_lahir = request.POST.get(
            'tempat_lahir',
            ''
        ).strip().title()

        tanggal_lahir = request.POST.get(
            'tanggal_lahir',
            ''
        ).strip()

        alamat = request.POST.get(
            'alamat',
            ''
        ).strip()

        no_tlpn_wa = request.POST.get(
            'no_tlpn_wa',
            ''
        ).strip()

        asrama_master = request.POST.get(
            'asrama_master'
        ) or None

        kelas = request.POST.get(
            'kelas'
        ) or None

        tahun_ajaran = request.POST.get(
            'tahun_ajaran'
        ) or None

        semester = request.POST.get(
            'semester',
            ''
        ).strip()

        prodi = request.POST.get(
            'prodi',
            ''
        ).strip()

        # -------------------------------------------------
        # VALIDASI
        # -------------------------------------------------

        if not nama:

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': 'Nama siswa wajib diisi.'
                }
            )

        if not tanggal_lahir:

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': 'Tanggal lahir wajib diisi.'
                }
            )

        if not nim:

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': 'NIM wajib diisi.'
                }
            )

        if not nim.isdigit():

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': 'NIM hanya boleh berisi angka.'
                }
            )

        if Student.objects.filter(
            nim=nim
        ).exists():

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': (
                        f'NIM {nim} sudah digunakan. '
                        'Silakan gunakan NIM lain.'
                    )
                }
            )

        if (
            no_tlpn_wa
            and not no_tlpn_wa.isdigit()
        ):

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': (
                        'No. Telp / WA hanya boleh '
                        'berisi angka.'
                    )
                }
            )

        if not tahun_ajaran:

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': 'Tahun ajaran wajib dipilih.'
                }
            )

        if not TahunAjaran.objects.filter(id=tahun_ajaran).exists():

            return render(
                request,
                'students/student_form.html',
                {
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'error': 'Tahun ajaran tidak ditemukan.'
                }
            )

        # -------------------------------------------------
        # SIMPAN SISWA
        # -------------------------------------------------

        Student.objects.create(
            nama=nama,
            nama_ayah=nama_ayah,
            nim=nim,
            jk=jk,
            tempat_lahir=tempat_lahir,
            tanggal_lahir=tanggal_lahir,
            alamat=alamat,
            no_tlpn_wa=no_tlpn_wa,
            asrama_master_id=asrama_master,
            kelas_id=kelas,
            tahun_ajaran_id=tahun_ajaran,
            semester=semester,
            prodi=prodi,
        )

        messages.success(
            request,
            f'Data siswa {nama} berhasil ditambahkan.'
        )

        return redirect('student_list')

    return render(
        request,
        'students/student_form.html',
        {
            'asramas': asramas,
            'kelass': kelass,
            'tahun_ajarans': tahun_ajarans,
        }
    )


# =========================================================
# EDIT SISWA
# =========================================================

def student_edit(request, id):

    student = get_object_or_404(
        Student,
        id=id
    )

    asramas = (
        Asrama.objects
        .all()
        .order_by('nama')
    )

    kelass = (
        Kelas.objects
        .all()
        .order_by('nama')
    )

    tahun_ajarans = (
        TahunAjaran.objects
        .all()
        .order_by('-nama')
    )


    if request.method == 'POST':

        student.nama = request.POST.get(
            'nama',
            ''
        ).strip().title()

        student.nama_ayah = request.POST.get(
            'nama_ayah',
            ''
        ).strip().title()

        student.nim = request.POST.get(
            'nim',
            ''
        ).strip()

        if not student.nim:

            return render(
                request,
                'students/student_form.html',
                {
                    'student': student,
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'edit_mode': True,
                    'error': 'NIM wajib diisi.'
                }
            )

        if not student.nim.isdigit():

            return render(
                request,
                'students/student_form.html',
                {
                    'student': student,
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'edit_mode': True,
                    'error': (
                        'NIM hanya boleh berisi angka.'
                    )
                }
            )

        if (
            Student.objects
            .filter(nim=student.nim)
            .exclude(id=student.id)
            .exists()
        ):

            return render(
                request,
                'students/student_form.html',
                {
                    'student': student,
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'edit_mode': True,
                    'error': (
                        f'NIM {student.nim} sudah '
                        'digunakan oleh siswa lain.'
                    )
                }
            )

        student.jk = request.POST.get(
            'jk',
            ''
        ).strip()

        student.tempat_lahir = request.POST.get(
            'tempat_lahir',
            ''
        ).strip().title()

        student.tanggal_lahir = request.POST.get(
            'tanggal_lahir',
            ''
        ).strip()

        student.alamat = request.POST.get(
            'alamat',
            ''
        ).strip()

        student.no_tlpn_wa = request.POST.get(
            'no_tlpn_wa',
            ''
        ).strip()

        if (
            student.no_tlpn_wa
            and not student.no_tlpn_wa.isdigit()
        ):

            return render(
                request,
                'students/student_form.html',
                {
                    'student': student,
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'edit_mode': True,
                    'error': (
                        'No. Telp / WA hanya boleh '
                        'berisi angka.'
                    )
                }
            )

        student.asrama_master_id = request.POST.get(
            'asrama_master'
        ) or None

        student.kelas_id = request.POST.get(
            'kelas'
        ) or None

        tahun_ajaran_id = request.POST.get(
            'tahun_ajaran'
        ) or None

        if not tahun_ajaran_id:
            return render(
                request,
                'students/student_form.html',
                {
                    'student': student,
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'edit_mode': True,
                    'error': 'Tahun ajaran wajib dipilih.'
                }
            )

        if not TahunAjaran.objects.filter(id=tahun_ajaran_id).exists():
            return render(
                request,
                'students/student_form.html',
                {
                    'student': student,
                    'asramas': asramas,
                    'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
                    'edit_mode': True,
                    'error': 'Tahun ajaran tidak ditemukan.'
                }
            )

        student.tahun_ajaran_id = tahun_ajaran_id

        student.semester = request.POST.get(
            'semester',
            ''
        ).strip()

        student.prodi = request.POST.get(
            'prodi',
            ''
        ).strip()

        student.status = (
            request.POST.get('status') == 'True'
        )

        student.save()

        messages.success(
            request,
            f'Data siswa {student.nama} berhasil diperbarui.'
        )

        return redirect('student_list')

    return render(
        request,
        'students/student_form.html',
        {
            'student': student,
            'asramas': asramas,
            'kelass': kelass,
                    'tahun_ajarans': tahun_ajarans,
            'edit_mode': True,
        }
    )


# =========================================================
# HAPUS SISWA
# =========================================================

def student_delete(request, id):

    student = get_object_or_404(
        Student,
        id=id
    )

    if request.method == 'POST':

        nama = student.nama

        student.delete()

        messages.success(
            request,
            f'Data siswa {nama} berhasil dihapus.'
        )

        return redirect('student_list')

    return render(
        request,
        'students/student_confirm_delete.html',
        {
            'student': student
        }
    )


# =========================================================
# ABSENSI SISWA
# =========================================================

def absensi(request):

    # =====================================================
    # POST - SIMPAN ABSENSI
    # =====================================================

    if request.method == 'POST':

        tanggal = request.POST.get('tanggal', '').strip()
        kelas_id = request.POST.get('kelas', '').strip()

        # ---------------------------------------------
        # Validasi tanggal
        # ---------------------------------------------

        try:
            tanggal_obj = datetime.strptime(
                tanggal,
                '%Y-%m-%d'
            ).date()

        except (ValueError, TypeError):

            tanggal_obj = timezone.localdate()
            tanggal = tanggal_obj.strftime('%Y-%m-%d')

        # ---------------------------------------------
        # Ambil siswa aktif
        # ---------------------------------------------

        students = (
            Student.objects
            .filter(status=True)
            .order_by('nama')
        )

        # Jika memilih kelas
        if kelas_id:
            students = students.filter(
                kelas_id=kelas_id
            )

        # ---------------------------------------------
        # Status yang diperbolehkan
        # ---------------------------------------------

        status_valid = [
            'Hadir',
            'Izin',
            'Sakit',
            'Alpa',
        ]

        # ---------------------------------------------
        # Simpan absensi setiap siswa
        # ---------------------------------------------

        for student in students:

            status = request.POST.get(
                f'status_{student.id}',
                'Alpa'
            )

            keterangan = request.POST.get(
                f'keterangan_{student.id}',
                ''
            ).strip()

            # Validasi status
            if status not in status_valid:
                status = 'Alpa'

            Absensi.objects.update_or_create(
                student=student,
                tanggal=tanggal_obj,
                defaults={
                    'status': status,
                    'keterangan': keterangan,
                }
            )

        messages.success(
            request,
            f'Presensi siswa untuk tanggal '
            f'{tanggal_obj.strftime("%d/%m/%Y")} berhasil disimpan.',
            extra_tags='presensi_siswa'
        )

        # ---------------------------------------------
        # KEMBALI KE HALAMAN ABSENSI
        # ---------------------------------------------

        return redirect(
            f"{reverse('presensi_siswa')}"
            f"?tanggal={tanggal}"
            f"&kelas={kelas_id}"
        )

    # =====================================================
    # GET - TAMPILKAN ABSENSI
    # =====================================================

    tanggal = request.GET.get(
        'tanggal',
        timezone.localdate().isoformat()
    )

    # Validasi tanggal GET
    try:

        tanggal_obj = datetime.strptime(
            tanggal,
            '%Y-%m-%d'
        ).date()

    except (ValueError, TypeError):

        tanggal_obj = timezone.localdate()
        tanggal = tanggal_obj.isoformat()

    # ---------------------------------------------
    # Filter kelas
    # ---------------------------------------------

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    # ---------------------------------------------
    # Ambil siswa aktif
    # ---------------------------------------------

    students = (
        Student.objects
        .select_related('kelas')
        .filter(status=True)
        .order_by('nama')
    )

    if kelas_id:
        students = students.filter(
            kelas_id=kelas_id
        )

    # ---------------------------------------------
    # Ambil data absensi tanggal tersebut
    # ---------------------------------------------

    absensi_data = (
        Absensi.objects
        .filter(
            tanggal=tanggal_obj,
            student__in=students
        )
    )

    # ---------------------------------------------
    # Buat map absensi
    # ---------------------------------------------

    absensi_map = {
        absensi.student_id: absensi
        for absensi in absensi_data
    }

    # ---------------------------------------------
    # Tempel data absensi ke masing-masing siswa
    # ---------------------------------------------

    for student in students:

        student.absensi_hari_ini = (
            absensi_map.get(student.id)
        )

    # ---------------------------------------------
    # Ambil semua kelas
    # ---------------------------------------------

    kelass = (
        Kelas.objects
        .all()
        .order_by('nama')
    )

    # ---------------------------------------------
    # Render
    # ---------------------------------------------

    context = {
        'students': students,
        'kelass': kelass,
        'tanggal': tanggal,
        'kelas_id': kelas_id,
    }

    return render(
        request,
        'students/absensi.html',
        context
    )

# =========================================================
# REKAP ABSENSI
# =========================================================

def rekap_absensi(request):

    # =====================================================
    # TANGGAL
    # =====================================================

    tanggal_hari_ini = timezone.localdate()

    tanggal_mulai = request.GET.get(
        'tanggal_mulai',
        ''
    ).strip()

    tanggal_sampai = request.GET.get(
        'tanggal_sampai',
        ''
    ).strip()

    # -----------------------------------------------------
    # TANGGAL MULAI
    # -----------------------------------------------------

    if not tanggal_mulai:

        tanggal_mulai_obj = tanggal_hari_ini.replace(day=1)

        tanggal_mulai = tanggal_mulai_obj.strftime('%Y-%m-%d')

    else:

        try:

            tanggal_mulai_obj = datetime.strptime(
                tanggal_mulai,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            tanggal_mulai_obj = tanggal_hari_ini.replace(day=1)

            tanggal_mulai = tanggal_mulai_obj.strftime('%Y-%m-%d')

    # -----------------------------------------------------
    # TANGGAL SAMPAI
    # -----------------------------------------------------

    if not tanggal_sampai:

        bulan_berikutnya = (
            tanggal_mulai_obj.replace(day=28)
            + timedelta(days=4)
        )

        tanggal_sampai_obj = (
            bulan_berikutnya
            - timedelta(days=bulan_berikutnya.day)
        )

        tanggal_sampai = tanggal_sampai_obj.strftime('%Y-%m-%d')

    else:

        try:

            tanggal_sampai_obj = datetime.strptime(
                tanggal_sampai,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            bulan_berikutnya = (
                tanggal_mulai_obj.replace(day=28)
                + timedelta(days=4)
            )

            tanggal_sampai_obj = (
                bulan_berikutnya
                - timedelta(days=bulan_berikutnya.day)
            )

            tanggal_sampai = tanggal_sampai_obj.strftime('%Y-%m-%d')

    # -----------------------------------------------------
    # JIKA TANGGAL TERBALIK
    # -----------------------------------------------------

    if tanggal_mulai_obj > tanggal_sampai_obj:

        tanggal_mulai_obj, tanggal_sampai_obj = (
            tanggal_sampai_obj,
            tanggal_mulai_obj
        )

        tanggal_mulai = tanggal_mulai_obj.strftime('%Y-%m-%d')
        tanggal_sampai = tanggal_sampai_obj.strftime('%Y-%m-%d')

    # =====================================================
    # DAFTAR TANGGAL
    # =====================================================

    jumlah_hari = (
        tanggal_sampai_obj - tanggal_mulai_obj
    ).days + 1

    tanggal_list = [
        tanggal_mulai_obj + timedelta(days=i)
        for i in range(jumlah_hari)
    ]

    hari_bulan = [
        tanggal.day
        for tanggal in tanggal_list
    ]

    # =====================================================
    # FILTER KELAS
    # =====================================================

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    kelass = (
        Kelas.objects
        .all()
        .order_by('nama')
    )

    # =====================================================
    # SISWA AKTIF
    # =====================================================

    students = (
        Student.objects
        .select_related('kelas')
        .filter(status=True)
        .order_by('nama')
    )

    if kelas_id:

        students = students.filter(
            kelas_id=kelas_id
        )

    # =====================================================
    # ABSENSI
    # =====================================================

    absensis = (
        Absensi.objects
        .select_related(
            'student',
            'student__kelas'
        )
        .filter(
            tanggal__range=[
                tanggal_mulai_obj,
                tanggal_sampai_obj
            ],
            student__in=students
        )
        .order_by('tanggal')
    )

    # =====================================================
    # MAP ABSENSI
    # =====================================================

    absensi_map = {}

    for absensi_data in absensis:

        if absensi_data.student_id not in absensi_map:

            absensi_map[absensi_data.student_id] = {}

        absensi_map[
            absensi_data.student_id
        ][
            absensi_data.tanggal
        ] = absensi_data.status

    # =====================================================
    # REKAP PER SISWA
    # =====================================================

    rekap_bulanan_siswa = []

    total_hadir = 0
    total_izin = 0
    total_sakit = 0
    total_alpa = 0

    for student in students:

        data_hari = absensi_map.get(
            student.id,
            {}
        )

        hadir = 0
        izin = 0
        sakit = 0
        alpa = 0

        # -------------------------------------------------
        # SIAPKAN STATUS SETIAP TANGGAL
        # -------------------------------------------------

        daftar_hari = []

        for tanggal in tanggal_list:

            status = data_hari.get(tanggal)

            if status == 'Hadir':
                hadir += 1

            elif status == 'Izin':
                izin += 1

            elif status == 'Sakit':
                sakit += 1

            elif status == 'Alpa':
                alpa += 1

            daftar_hari.append({
                'tanggal': tanggal,
                'hari': tanggal.day,
                'status': status,
            })

        # -------------------------------------------------
        # TOTAL
        # -------------------------------------------------

        total = (
            hadir
            + izin
            + sakit
            + alpa
        )

        if total > 0:

            persentase = round(
                (hadir / total) * 100,
                1
            )

        else:

            persentase = 0

        rekap_bulanan_siswa.append({
            'student': student,
            'hari': data_hari,
            'daftar_hari': daftar_hari,
            'hadir': hadir,
            'izin': izin,
            'sakit': sakit,
            'alpa': alpa,
            'total': total,
            'persentase': persentase,
        })

        total_hadir += hadir
        total_izin += izin
        total_sakit += sakit
        total_alpa += alpa

    # =====================================================
    # STATISTIK
    # =====================================================

    total_absensi = (
        total_hadir
        + total_izin
        + total_sakit
        + total_alpa
    )

    if total_absensi > 0:

        persentase_kehadiran = round(
            (
                total_hadir
                / total_absensi
            ) * 100,
            1
        )

    else:

        persentase_kehadiran = 0

    # =====================================================
    # NAMA BULAN
    # =====================================================

    nama_bulan_indonesia = {
        'January': 'Januari',
        'February': 'Februari',
        'March': 'Maret',
        'April': 'April',
        'May': 'Mei',
        'June': 'Juni',
        'July': 'Juli',
        'August': 'Agustus',
        'September': 'September',
        'October': 'Oktober',
        'November': 'November',
        'December': 'Desember',
    }

    if (
        tanggal_mulai_obj.month == tanggal_sampai_obj.month
        and
        tanggal_mulai_obj.year == tanggal_sampai_obj.year
    ):

        nama_bulan = (
            nama_bulan_indonesia.get(
                tanggal_mulai_obj.strftime('%B'),
                tanggal_mulai_obj.strftime('%B')
            )
            + f' {tanggal_mulai_obj.year}'
        )

    else:

        nama_bulan = (
            tanggal_mulai_obj.strftime('%d/%m/%Y')
            + ' - '
            + tanggal_sampai_obj.strftime('%d/%m/%Y')
        )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'bulan':
            tanggal_mulai_obj.strftime('%Y-%m'),

        'nama_bulan':
            nama_bulan,

        'tahun':
            tanggal_mulai_obj.year,

        'nomor_bulan':
            tanggal_mulai_obj.month,

        'tanggal_list':
            tanggal_list,

        'hari_bulan':
            hari_bulan,

        'jumlah_hari':
            jumlah_hari,

        'tanggal_mulai':
            tanggal_mulai_obj,

        'tanggal_sampai':
            tanggal_sampai_obj,

        'tanggal_mulai_obj':
            tanggal_mulai_obj,

        'tanggal_sampai_obj':
            tanggal_sampai_obj,

        'tanggal_mulai_value':
            tanggal_mulai,

        'tanggal_sampai_value':
            tanggal_sampai,

        'kelas_id':
            kelas_id,

        'kelass':
            kelass,

        'rekap_bulanan_siswa':
            rekap_bulanan_siswa,

        'total_hadir':
            total_hadir,

        'total_izin':
            total_izin,

        'total_sakit':
            total_sakit,

        'total_alpa':
            total_alpa,

        'total_absensi':
            total_absensi,

        'persentase_kehadiran':
            persentase_kehadiran,

        'hari_efektif':
            jumlah_hari,
    }

    return render(
        request,
        'students/rekap_absensi.html',
        context
    )


# =========================================================
# INPUT PRESENSI GURU
# =========================================================

def presensi_guru(request):

    # =====================================================
    # POST - SIMPAN PRESENSI
    # =====================================================

    if request.method == 'POST':

        tanggal = request.POST.get(
            'tanggal',
            ''
        ).strip()

        # -------------------------------------------------
        # VALIDASI TANGGAL
        # -------------------------------------------------

        if not tanggal:

            messages.error(
                request,
                'Tanggal presensi wajib dipilih.'
            )

            return redirect('presensi_guru')

        try:

            tanggal_obj = datetime.strptime(
                tanggal,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            messages.error(
                request,
                'Format tanggal tidak valid.'
            )

            return redirect('presensi_guru')

        # -------------------------------------------------
        # GURU
        # -------------------------------------------------

        gurus = (
            Guru.objects
            .all()
            .order_by('nama')
        )

        # -------------------------------------------------
        # STATUS VALID
        # -------------------------------------------------

        status_valid = [
            'Hadir',
            'Izin',
            'Sakit',
            'Alpa',
        ]

        jumlah_disimpan = 0

        # -------------------------------------------------
        # SIMPAN DATA
        # -------------------------------------------------

        for guru in gurus:

            status = request.POST.get(
                f'status_{guru.id}',
                'Alpa'
            )

            keterangan = request.POST.get(
                f'keterangan_{guru.id}',
                ''
            ).strip()

            if status not in status_valid:
                status = 'Alpa'

            PresensiGuru.objects.update_or_create(

                guru=guru,

                tanggal=tanggal_obj,

                defaults={
                    'status': status,
                    'keterangan': keterangan,
                }

            )

            jumlah_disimpan += 1

        # -------------------------------------------------
        # PESAN
        # -------------------------------------------------

        messages.success(
            request,
            f'Presensi guru untuk tanggal '
            f'{tanggal_obj.strftime("%d/%m/%Y")} berhasil disimpan.',
            extra_tags='presensi_guru'
        )

        return redirect(
            f"{reverse('presensi_guru')}"
            f"?tanggal={tanggal}"
        )

    # =====================================================
    # GET - TAMPILKAN DATA
    # =====================================================

    tanggal = request.GET.get(
        'tanggal',
        ''
    ).strip()

    if not tanggal:

        tanggal = (
            timezone
            .localdate()
            .isoformat()
        )

    # -----------------------------------------------------
    # VALIDASI TANGGAL
    # -----------------------------------------------------

    try:

        tanggal_obj = datetime.strptime(
            tanggal,
            '%Y-%m-%d'
        ).date()

        tanggal = tanggal_obj.isoformat()

    except ValueError:

        tanggal_obj = timezone.localdate()

        tanggal = tanggal_obj.isoformat()

    # -----------------------------------------------------
    # GURU
    # -----------------------------------------------------

    gurus = (
        Guru.objects
        .all()
        .order_by('nama')
    )

    # -----------------------------------------------------
    # PRESENSI PADA TANGGAL TERPILIH
    # -----------------------------------------------------

    presensi_data = (
        PresensiGuru.objects
        .filter(
            tanggal=tanggal_obj,
            guru__in=gurus
        )
    )

    presensi_map = {
        presensi.guru_id: presensi
        for presensi in presensi_data
    }

    # -----------------------------------------------------
    # TEMPELKAN PRESENSI KE GURU
    # -----------------------------------------------------

    for guru in gurus:

        guru.presensi_hari_ini = (
            presensi_map.get(guru.id)
        )

    # -----------------------------------------------------
    # RENDER
    # -----------------------------------------------------

    return render(
        request,
        'students/presensi_guru.html',
        {
            'gurus': gurus,
            'tanggal': tanggal,
        }
    )
# =========================================================
# REKAP PRESENSI GURU
# =========================================================
def rekap_presensi_guru(request):

    # =====================================================
    # TANGGAL
    # =====================================================

    tanggal_hari_ini = timezone.localdate()

    tanggal_mulai = request.GET.get('tanggal_mulai', '').strip()
    tanggal_sampai = request.GET.get('tanggal_sampai', '').strip()

    # -------------------------
    # TANGGAL MULAI
    # -------------------------

    if not tanggal_mulai:

        tanggal_mulai_obj = tanggal_hari_ini.replace(day=1)

        tanggal_mulai = tanggal_mulai_obj.strftime('%Y-%m-%d')

    else:

        try:

            tanggal_mulai_obj = datetime.strptime(
                tanggal_mulai,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            tanggal_mulai_obj = tanggal_hari_ini.replace(day=1)

            tanggal_mulai = tanggal_mulai_obj.strftime('%Y-%m-%d')


    # -------------------------
    # TANGGAL SAMPAI
    # -------------------------

    if not tanggal_sampai:

        bulan_berikutnya = (
            tanggal_mulai_obj.replace(day=28)
            + timedelta(days=4)
        )

        tanggal_sampai_obj = (
            bulan_berikutnya
            - timedelta(days=bulan_berikutnya.day)
        )

        tanggal_sampai = tanggal_sampai_obj.strftime('%Y-%m-%d')

    else:

        try:

            tanggal_sampai_obj = datetime.strptime(
                tanggal_sampai,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            bulan_berikutnya = (
                tanggal_mulai_obj.replace(day=28)
                + timedelta(days=4)
            )

            tanggal_sampai_obj = (
                bulan_berikutnya
                - timedelta(days=bulan_berikutnya.day)
            )

            tanggal_sampai = tanggal_sampai_obj.strftime('%Y-%m-%d')


    # =====================================================
    # JIKA TANGGAL TERBALIK
    # =====================================================

    if tanggal_mulai_obj > tanggal_sampai_obj:

        tanggal_mulai_obj, tanggal_sampai_obj = (
            tanggal_sampai_obj,
            tanggal_mulai_obj
        )

        tanggal_mulai = tanggal_mulai_obj.strftime('%Y-%m-%d')

        tanggal_sampai = tanggal_sampai_obj.strftime('%Y-%m-%d')


    # =====================================================
    # DAFTAR TANGGAL
    # =====================================================

    jumlah_hari = (
        tanggal_sampai_obj - tanggal_mulai_obj
    ).days + 1

    tanggal_list = [
        tanggal_mulai_obj + timedelta(days=i)
        for i in range(jumlah_hari)
    ]

    hari_bulan = [
        tanggal.day
        for tanggal in tanggal_list
    ]


    # =====================================================
    # DATA GURU
    # =====================================================

    gurus = (
        Guru.objects
        .all()
        .order_by('nama')
    )


    # =====================================================
    # DATA PRESENSI GURU
    # =====================================================

    presensis = (
        PresensiGuru.objects
        .select_related('guru')
        .filter(
            tanggal__range=[
                tanggal_mulai_obj,
                tanggal_sampai_obj
            ],
            guru__in=gurus
        )
        .order_by(
            'tanggal',
            'guru__nama'
        )
    )


    # =====================================================
    # BUAT MAP PRESENSI
    # =====================================================

    presensi_map = {}

    for presensi in presensis:

        if presensi.guru_id not in presensi_map:

            presensi_map[presensi.guru_id] = {}

        presensi_map[
            presensi.guru_id
        ][
            presensi.tanggal
        ] = presensi.status


    # =====================================================
    # REKAP
    # =====================================================

    rekap_bulanan_guru = []

    total_hadir = 0
    total_izin = 0
    total_sakit = 0
    total_alpa = 0


    # =====================================================
    # LOOP GURU
    # =====================================================

    for guru in gurus:

        data_hari = presensi_map.get(
            guru.id,
            {}
        )

        hadir = 0
        izin = 0
        sakit = 0
        alpa = 0

        daftar_hari = []


        # ---------------------------------------------
        # LOOP SETIAP TANGGAL
        # ---------------------------------------------

        for tanggal in tanggal_list:

            status = data_hari.get(
                tanggal
            )


            # Simpan data untuk template
            daftar_hari.append({

                'tanggal': tanggal,

                'status': status,

            })


            # Hitung statistik

            if status == 'Hadir':

                hadir += 1

            elif status == 'Izin':

                izin += 1

            elif status == 'Sakit':

                sakit += 1

            elif status == 'Alpa':

                alpa += 1


        # =================================================
        # TOTAL
        # =================================================

        total = (
            hadir
            + izin
            + sakit
            + alpa
        )


        # =================================================
        # PERSENTASE
        # =================================================

        if total > 0:

            persentase = round(
                (hadir / total) * 100,
                1
            )

        else:

            persentase = 0


        # =================================================
        # MASUKKAN KE REKAP
        # =================================================

        rekap_bulanan_guru.append({

            'guru': guru,

            'hari': data_hari,

            'daftar_hari': daftar_hari,

            'hadir': hadir,

            'izin': izin,

            'sakit': sakit,

            'alpa': alpa,

            'total': total,

            'persentase': persentase,

        })


        # =================================================
        # TOTAL SEMUA GURU
        # =================================================

        total_hadir += hadir

        total_izin += izin

        total_sakit += sakit

        total_alpa += alpa


    # =====================================================
    # TOTAL ABSENSI
    # =====================================================

    total_absensi = (
        total_hadir
        + total_izin
        + total_sakit
        + total_alpa
    )


    # =====================================================
    # PERSENTASE KEHADIRAN
    # =====================================================

    if total_absensi > 0:

        persentase_kehadiran = round(
            (total_hadir / total_absensi) * 100,
            1
        )

    else:

        persentase_kehadiran = 0


    # =====================================================
    # NAMA BULAN INDONESIA
    # =====================================================

    nama_bulan_indonesia = {

        'January': 'Januari',

        'February': 'Februari',

        'March': 'Maret',

        'April': 'April',

        'May': 'Mei',

        'June': 'Juni',

        'July': 'Juli',

        'August': 'Agustus',

        'September': 'September',

        'October': 'Oktober',

        'November': 'November',

        'December': 'Desember',

    }


    # =====================================================
    # NAMA PERIODE
    # =====================================================

    if (
        tanggal_mulai_obj.month
        == tanggal_sampai_obj.month

        and

        tanggal_mulai_obj.year
        == tanggal_sampai_obj.year
    ):

        nama_bulan = (

            nama_bulan_indonesia.get(

                tanggal_mulai_obj.strftime('%B'),

                tanggal_mulai_obj.strftime('%B')

            )

            + f' {tanggal_mulai_obj.year}'

        )

    else:

        nama_bulan = (

            tanggal_mulai_obj.strftime('%d/%m/%Y')

            + ' - '

            + tanggal_sampai_obj.strftime('%d/%m/%Y')

        )


    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'bulan':
            tanggal_mulai_obj.strftime('%Y-%m'),

        'nama_bulan':
            nama_bulan,

        'tahun':
            tanggal_mulai_obj.year,

        'nomor_bulan':
            tanggal_mulai_obj.month,

        'tanggal_list':
            tanggal_list,

        'hari_bulan':
            hari_bulan,

        'jumlah_hari':
            jumlah_hari,

        'tanggal_mulai':
            tanggal_mulai_obj,

        'tanggal_sampai':
            tanggal_sampai_obj,

        'tanggal_mulai_obj':
            tanggal_mulai_obj,

        'tanggal_sampai_obj':
            tanggal_sampai_obj,

        'tanggal_mulai_value':
            tanggal_mulai,

        'tanggal_sampai_value':
            tanggal_sampai,

        'rekap_bulanan_guru':
            rekap_bulanan_guru,

        'total_hadir':
            total_hadir,

        'total_izin':
            total_izin,

        'total_sakit':
            total_sakit,

        'total_alpa':
            total_alpa,

        'total_absensi':
            total_absensi,

        'persentase_kehadiran':
            persentase_kehadiran,

        'hari_efektif':
            jumlah_hari,

    }


    return render(
        request,
        'students/rekap_presensi_guru.html',
        context
    )

def cetak_rekap_presensi_guru(request):

    # =====================================================
    # FILTER TANGGAL
    # =====================================================

    tanggal_mulai = request.GET.get('tanggal_mulai')
    tanggal_sampai = request.GET.get('tanggal_sampai')

    # Default: bulan berjalan
    if not tanggal_mulai:
        tanggal_mulai_obj = date.today().replace(day=1)
    else:
        tanggal_mulai_obj = datetime.strptime(
            tanggal_mulai,
            '%Y-%m-%d'
        ).date()

    if not tanggal_sampai:
        if tanggal_mulai:
            tanggal_sampai_obj = tanggal_mulai_obj.replace(
                day=calendar.monthrange(
                    tanggal_mulai_obj.year,
                    tanggal_mulai_obj.month
                )[1]
            )
        else:
            tanggal_sampai_obj = tanggal_mulai_obj.replace(
                day=calendar.monthrange(
                    tanggal_mulai_obj.year,
                    tanggal_mulai_obj.month
                )[1]
            )
    else:
        tanggal_sampai_obj = datetime.strptime(
            tanggal_sampai,
            '%Y-%m-%d'
        ).date()

    # Jika terbalik, tukar
    if tanggal_mulai_obj > tanggal_sampai_obj:
        tanggal_mulai_obj, tanggal_sampai_obj = (
            tanggal_sampai_obj,
            tanggal_mulai_obj
        )

    # =====================================================
    # DAFTAR TANGGAL
    # =====================================================

    tanggal_list = []

    tanggal_sekarang = tanggal_mulai_obj

    while tanggal_sekarang <= tanggal_sampai_obj:
        tanggal_list.append(tanggal_sekarang)
        tanggal_sekarang += timedelta(days=1)

    hari_bulan = [
        tanggal.day
        for tanggal in tanggal_list
    ]

    jumlah_hari = len(tanggal_list)

    # =====================================================
    # DATA GURU
    # =====================================================

    gurus = Guru.objects.all().order_by('nama')

    # =====================================================
    # DATA PRESENSI
    # =====================================================

    presensi_gurus = PresensiGuru.objects.filter(
        tanggal__range=[
            tanggal_mulai_obj,
            tanggal_sampai_obj
        ],
        guru__in=gurus
    ).select_related(
        'guru'
    ).order_by(
        'tanggal',
        'guru__nama'
    )

    # =====================================================
    # BUAT MAP PRESENSI
    # =====================================================

    presensi_map = {}

    for presensi in presensi_gurus:

        guru_id = presensi.guru_id

        if guru_id not in presensi_map:
            presensi_map[guru_id] = {}

        presensi_map[guru_id][
            presensi.tanggal
        ] = presensi.status

    # =====================================================
    # REKAP GURU
    # =====================================================

    rekap_bulanan_guru = []

    total_hadir = 0
    total_izin = 0
    total_sakit = 0
    total_alpa = 0

    for guru in gurus:

        daftar_hari = []

        guru_hadir = 0
        guru_izin = 0
        guru_sakit = 0
        guru_alpa = 0

        for tanggal in tanggal_list:

            status = presensi_map.get(
                guru.id,
                {}
            ).get(
                tanggal
            )

            daftar_hari.append({
                'tanggal': tanggal,
                'status': status
            })

            if status == 'Hadir':
                guru_hadir += 1

            elif status == 'Izin':
                guru_izin += 1

            elif status == 'Sakit':
                guru_sakit += 1

            elif status == 'Alpa':
                guru_alpa += 1

        total_absensi_guru = (
            guru_hadir
            + guru_izin
            + guru_sakit
            + guru_alpa
        )

        if total_absensi_guru > 0:
            persentase = round(
                (guru_hadir / total_absensi_guru) * 100,
                2
            )
        else:
            persentase = 0

        rekap_bulanan_guru.append({
            'guru': guru,
            'hari': daftar_hari,
            'hadir': guru_hadir,
            'izin': guru_izin,
            'sakit': guru_sakit,
            'alpa': guru_alpa,
            'total': total_absensi_guru,
            'persentase': persentase,
        })

        total_hadir += guru_hadir
        total_izin += guru_izin
        total_sakit += guru_sakit
        total_alpa += guru_alpa

    # =====================================================
    # TOTAL
    # =====================================================

    total_absensi = (
        total_hadir
        + total_izin
        + total_sakit
        + total_alpa
    )

    if total_absensi > 0:
        persentase_kehadiran = round(
            (total_hadir / total_absensi) * 100,
            2
        )
    else:
        persentase_kehadiran = 0

    # =====================================================
    # NAMA BULAN
    # =====================================================

    nama_bulan = [
        '',
        'Januari',
        'Februari',
        'Maret',
        'April',
        'Mei',
        'Juni',
        'Juli',
        'Agustus',
        'September',
        'Oktober',
        'November',
        'Desember'
    ]

    bulan = tanggal_mulai_obj.month
    tahun = tanggal_mulai_obj.year

    context = {
        'bulan': bulan,
        'nama_bulan': nama_bulan[bulan],
        'tahun': tahun,

        'tanggal_list': tanggal_list,
        'hari_bulan': hari_bulan,
        'jumlah_hari': jumlah_hari,

        'tanggal_mulai_obj': tanggal_mulai_obj,
        'tanggal_sampai_obj': tanggal_sampai_obj,

        'rekap_bulanan_guru': rekap_bulanan_guru,

        'total_hadir': total_hadir,
        'total_izin': total_izin,
        'total_sakit': total_sakit,
        'total_alpa': total_alpa,
        'total_absensi': total_absensi,
        'persentase_kehadiran': persentase_kehadiran,
    }

    return render(
        request,
        'students/cetak_rekap_presensi_guru.html',
        context
    )

# =========================================================
# PENILAIAN
# =========================================================

def penilaian(request):

    # =====================================================
    # DATA MASTER
    # =====================================================

    tahun_ajarans = (
        TahunAjaran.objects
        .all()
        .order_by('-nama')
    )

    kelass = (
        Kelas.objects
        .all()
        .order_by('nama')
    )

    # =====================================================
    # FILTER
    # =====================================================

    tahun_ajaran_id = request.GET.get(
        'tahun_ajaran',
        ''
    ).strip()

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    search = request.GET.get(
        'search',
        ''
    ).strip()

    # =====================================================
    # DEFAULT
    # =====================================================

    mata_pelajarans = MataPelajaran.objects.none()

    students = Student.objects.none()

    # =====================================================
    # GET DATA
    # =====================================================

    if tahun_ajaran_id:

        mata_pelajarans = (
            MataPelajaran.objects
            .filter(
                tahun_ajaran_id=tahun_ajaran_id,
                aktif=True
            )
            .prefetch_related('kelas')
            .order_by('nama')
        )

        if kelas_id:

            # -------------------------------------------------
            # MATA PELAJARAN SESUAI KELAS
            # -------------------------------------------------

            mata_pelajarans = (
                mata_pelajarans
                .filter(
                    kelas__id=kelas_id
                )
                .distinct()
            )

            # -------------------------------------------------
            # SISWA SESUAI TAHUN AJARAN + KELAS
            # -------------------------------------------------

            students = (
                Student.objects
                .select_related(
                    'kelas',
                    'tahun_ajaran'
                )
                .filter(
                    status=True,
                    kelas_id=kelas_id,
                    tahun_ajaran_id=tahun_ajaran_id,
                )
                .order_by('nama')
            )

            # -------------------------------------------------
            # PENCARIAN NAMA / NIM
            # -------------------------------------------------

            if search:

                students = students.filter(
                    Q(nama__icontains=search)
                    |
                    Q(nim__icontains=search)
                )

    # =====================================================
    # POST - SIMPAN SEMUA NILAI
    # =====================================================

    if request.method == 'POST':

        # -------------------------------------------------
        # AMBIL FILTER DARI FORM
        # -------------------------------------------------

        tahun_ajaran_id = (
            request.POST.get(
                'tahun_ajaran'
            )
            or ''
        ).strip()

        kelas_id = (
            request.POST.get(
                'kelas'
            )
            or ''
        ).strip()

        search = (
            request.POST.get(
                'search'
            )
            or ''
        ).strip()

        # -------------------------------------------------
        # SISWA
        # -------------------------------------------------

        students = (
            Student.objects
            .filter(
                status=True,
                kelas_id=kelas_id,
                tahun_ajaran_id=tahun_ajaran_id,
            )
            .order_by('nama')
        )

        # -------------------------------------------------
        # MATA PELAJARAN
        # -------------------------------------------------

        mata_pelajarans = (
            MataPelajaran.objects
            .filter(
                tahun_ajaran_id=tahun_ajaran_id,
                aktif=True,
                kelas__id=kelas_id
            )
            .distinct()
            .order_by('nama')
        )

        # -------------------------------------------------
        # SIMPAN NILAI
        # -------------------------------------------------

        for student in students:

            for mapel in mata_pelajarans:

                # -----------------------------------------
                # NILAI HARIAN
                # -----------------------------------------

                nilai_harian_raw = request.POST.get(
                    f'nilai_harian_{student.id}_{mapel.id}'
                )

                # -----------------------------------------
                # NILAI UJIAN
                # -----------------------------------------

                nilai_ujian_raw = request.POST.get(
                    f'nilai_ujian_{student.id}_{mapel.id}'
                )

                nilai_harian = None

                nilai_ujian = None

                # -----------------------------------------
                # VALIDASI NILAI HARIAN
                # -----------------------------------------

                if nilai_harian_raw not in [
                    None,
                    ''
                ]:

                    try:

                        nilai_harian = int(
                            nilai_harian_raw
                        )

                        if (
                            nilai_harian < 0
                            or nilai_harian > 100
                        ):

                            messages.error(
                                request,
                                f'Nilai harian '
                                f'{student.nama} - '
                                f'{mapel.nama} '
                                f'harus 0 sampai 100.'
                            )

                            continue

                    except ValueError:

                        messages.error(
                            request,
                            f'Nilai harian '
                            f'{student.nama} - '
                            f'{mapel.nama} '
                            f'tidak valid.'
                        )

                        continue

                # -----------------------------------------
                # VALIDASI NILAI UJIAN
                # -----------------------------------------

                if nilai_ujian_raw not in [
                    None,
                    ''
                ]:

                    try:

                        nilai_ujian = int(
                            nilai_ujian_raw
                        )

                        if (
                            nilai_ujian < 0
                            or nilai_ujian > 100
                        ):

                            messages.error(
                                request,
                                f'Nilai ujian '
                                f'{student.nama} - '
                                f'{mapel.nama} '
                                f'harus 0 sampai 100.'
                            )

                            continue

                    except ValueError:

                        messages.error(
                            request,
                            f'Nilai ujian '
                            f'{student.nama} - '
                            f'{mapel.nama} '
                            f'tidak valid.'
                        )

                        continue

                # -----------------------------------------
                # JIKA KOSONG -> HAPUS
                # -----------------------------------------

                if (
                    nilai_harian is None
                    and
                    nilai_ujian is None
                ):

                    Penilaian.objects.filter(
                        student=student,
                        mata_pelajaran=mapel
                    ).delete()

                    continue

                # -----------------------------------------
                # HITUNG NILAI AKHIR
                # -----------------------------------------

                if (
                    nilai_harian is not None
                    and
                    nilai_ujian is not None
                ):

                    nilai_akhir = round(
                        (
                            nilai_harian
                            +
                            nilai_ujian
                        ) / 2
                    )

                else:

                    nilai_akhir = None

                # -----------------------------------------
                # SIMPAN / UPDATE
                # -----------------------------------------

                Penilaian.objects.update_or_create(
                    student=student,
                    mata_pelajaran=mapel,
                    defaults={
                        'nilai_harian':
                            nilai_harian,

                        'nilai_ujian':
                            nilai_ujian,

                        'nilai_akhir':
                            nilai_akhir,
                    }
                )

        messages.success(
            request,
            'Semua nilai berhasil disimpan.'
        )

        # -------------------------------------------------
        # KEMBALI KE HALAMAN PENILAIAN
        # -------------------------------------------------

        return redirect(
            f'/penilaian/'
            f'?tahun_ajaran={tahun_ajaran_id}'
            f'&kelas={kelas_id}'
            f'&search={search}'
        )

    # =====================================================
    # AMBIL NILAI YANG SUDAH TERSIMPAN
    # =====================================================

    penilaians = (
        Penilaian.objects
        .filter(
            student__in=students,
            mata_pelajaran__in=mata_pelajarans
        )
        .select_related(
            'student',
            'mata_pelajaran'
        )
    )

    # =====================================================
    # BENTUK DICTIONARY
    #
    # penilaian_data[
    #     student_id
    # ][
    #     mapel_id
    # ] = nilai
    # =====================================================

    penilaian_data = {}

    for nilai in penilaians:

        if nilai.student_id not in penilaian_data:

            penilaian_data[
                nilai.student_id
            ] = {}

        penilaian_data[
            nilai.student_id
        ][
            nilai.mata_pelajaran_id
        ] = nilai

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'tahun_ajarans':
            tahun_ajarans,

        'kelass':
            kelass,

        'students':
            students,

        'mata_pelajarans':
            mata_pelajarans,

        'penilaian_data':
            penilaian_data,

        'tahun_ajaran_id':
            tahun_ajaran_id,

        'kelas_id':
            kelas_id,

        'search':
            search,
    }

    return render(
        request,
        'students/penilaian.html',
        context
    )

# =========================================================
# CETAK REKAP ABSENSI
# =========================================================

def cetak_rekap_absensi(request):

    tanggal_mulai = request.GET.get(
        'tanggal_mulai',
        ''
    ).strip()

    tanggal_sampai = request.GET.get(
        'tanggal_sampai',
        ''
    ).strip()

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    today = timezone.localdate()

    # =====================================================
    # TANGGAL MULAI
    # =====================================================

    if tanggal_mulai:

        try:

            tanggal_mulai_obj = datetime.strptime(
                tanggal_mulai,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            tanggal_mulai_obj = today.replace(day=1)

    else:

        tanggal_mulai_obj = today.replace(day=1)

    # =====================================================
    # TANGGAL SAMPAI
    # =====================================================

    if tanggal_sampai:

        try:

            tanggal_sampai_obj = datetime.strptime(
                tanggal_sampai,
                '%Y-%m-%d'
            ).date()

        except ValueError:

            bulan_berikutnya = (
                tanggal_mulai_obj.replace(day=28)
                + timedelta(days=4)
            )

            tanggal_sampai_obj = (
                bulan_berikutnya
                - timedelta(
                    days=bulan_berikutnya.day
                )
            )

    else:

        bulan_berikutnya = (
            tanggal_mulai_obj.replace(day=28)
            + timedelta(days=4)
        )

        tanggal_sampai_obj = (
            bulan_berikutnya
            - timedelta(
                days=bulan_berikutnya.day
            )
        )

    # =====================================================
    # JIKA TANGGAL TERBALIK
    # =====================================================

    if tanggal_mulai_obj > tanggal_sampai_obj:

        tanggal_mulai_obj, tanggal_sampai_obj = (
            tanggal_sampai_obj,
            tanggal_mulai_obj
        )

    # =====================================================
    # SISWA AKTIF
    # =====================================================

    students = (
        Student.objects
        .select_related('kelas')
        .filter(status=True)
        .order_by('nama')
    )

    # =====================================================
    # FILTER KELAS
    # =====================================================

    if kelas_id:

        students = students.filter(
            kelas_id=kelas_id
        )

    # =====================================================
    # DATA ABSENSI
    # =====================================================

    absensis = (
        Absensi.objects
        .select_related(
            'student',
            'student__kelas'
        )
        .filter(
            tanggal__range=[
                tanggal_mulai_obj,
                tanggal_sampai_obj
            ],
            student__in=students
        )
        .order_by(
            'tanggal',
            'student__nama'
        )
    )

    # =====================================================
    # REKAP SISWA
    # =====================================================

    rekap_bulanan_siswa = []

    total_hadir = 0
    total_izin = 0
    total_sakit = 0
    total_alpa = 0

    for student in students:

        data_siswa = absensis.filter(
            student=student
        )

        hadir = data_siswa.filter(
            status='Hadir'
        ).count()

        alpa = data_siswa.filter(
            status='Alpa'
        ).count()

        izin = data_siswa.filter(
            status='Izin'
        ).count()

        sakit = data_siswa.filter(
            status='Sakit'
        ).count()

        total = (
            hadir
            + alpa
            + izin
            + sakit
        )

        if total > 0:

            persentase = round(
                (
                    hadir / total
                ) * 100,
                1
            )

        else:

            persentase = 0

        rekap_bulanan_siswa.append(
            {
                'student': student,
                'hadir': hadir,
                'alpa': alpa,
                'izin': izin,
                'sakit': sakit,
                'total': total,
                'persentase': persentase,
            }
        )

        total_hadir += hadir
        total_alpa += alpa
        total_izin += izin
        total_sakit += sakit

    # =====================================================
    # TOTAL ABSENSI
    # =====================================================

    total_absensi = (
        total_hadir
        + total_alpa
        + total_izin
        + total_sakit
    )

    # =====================================================
    # PERSENTASE
    # =====================================================

    if total_absensi > 0:

        persentase_kehadiran = round(
            (
                total_hadir
                / total_absensi
            ) * 100,
            1
        )

    else:

        persentase_kehadiran = 0

    # =====================================================
    # JUMLAH HARI
    # =====================================================

    jumlah_hari = (
        tanggal_sampai_obj
        - tanggal_mulai_obj
    ).days + 1

    # =====================================================
    # HARI EFEKTIF
    # =====================================================

    hari_efektif = (
        absensis
        .values('tanggal')
        .distinct()
        .count()
    )

    # =====================================================
    # KELAS
    # =====================================================

    kelas = None

    if kelas_id:

        kelas = (
            Kelas.objects
            .filter(id=kelas_id)
            .first()
        )

    # =====================================================
    # NAMA BULAN
    # =====================================================

    nama_bulan_indonesia = {

        'January': 'Januari',
        'February': 'Februari',
        'March': 'Maret',
        'April': 'April',
        'May': 'Mei',
        'June': 'Juni',
        'July': 'Juli',
        'August': 'Agustus',
        'September': 'September',
        'October': 'Oktober',
        'November': 'November',
        'December': 'Desember',

    }

    if (
        tanggal_mulai_obj.month
        == tanggal_sampai_obj.month
        and
        tanggal_mulai_obj.year
        == tanggal_sampai_obj.year
    ):

        nama_bulan = (
            nama_bulan_indonesia.get(
                tanggal_mulai_obj.strftime('%B'),
                tanggal_mulai_obj.strftime('%B')
            )
            + f' {tanggal_mulai_obj.year}'
        )

    else:

        nama_bulan = (
            tanggal_mulai_obj.strftime('%d/%m/%Y')
            + ' - '
            + tanggal_sampai_obj.strftime('%d/%m/%Y')
        )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'bulan':
            tanggal_mulai_obj.strftime('%Y-%m'),

        'nama_bulan':
            nama_bulan,

        'tahun':
            tanggal_mulai_obj.year,

        'nomor_bulan':
            tanggal_mulai_obj.month,

        'tanggal_list':
            [
                tanggal_mulai_obj
                + timedelta(days=i)
                for i in range(jumlah_hari)
            ],

        'hari_bulan':
            [
                (
                    tanggal_mulai_obj
                    + timedelta(days=i)
                ).day
                for i in range(jumlah_hari)
            ],

        'jumlah_hari':
            jumlah_hari,

        'tanggal_mulai':
            tanggal_mulai_obj,

        'tanggal_sampai':
            tanggal_sampai_obj,

        'tanggal_mulai_obj':
            tanggal_mulai_obj,

        'tanggal_sampai_obj':
            tanggal_sampai_obj,

        'tanggal_mulai_value':
            tanggal_mulai_obj.strftime('%Y-%m-%d'),

        'tanggal_sampai_value':
            tanggal_sampai_obj.strftime('%Y-%m-%d'),

        'kelas_id':
            kelas_id,

        'kelas':
            kelas,

        'kelas_obj':
            kelas,

        'rekap_bulanan_siswa':
            rekap_bulanan_siswa,

        'rekap_siswa':
            rekap_bulanan_siswa,

        'absensis':
            absensis,

        'total_hadir':
            total_hadir,

        'total_izin':
            total_izin,

        'total_sakit':
            total_sakit,

        'total_alpa':
            total_alpa,

        'total_absensi':
            total_absensi,

        'persentase_kehadiran':
            persentase_kehadiran,

        'hari_efektif':
            hari_efektif,

        'today':
            today,
    }

    return render(
        request,
        'students/cetak_rekap_absensi.html',
        context
    )


def data_kelas(request):
    kelass = Kelas.objects.all().order_by('nama')

    context = {
        'kelass': kelass,
    }

    return render(request, 'students/data_kelas.html', context)

def tambah_kelas(request):
    if request.method == 'POST':
        nama = request.POST.get('nama', '').strip()

        if not nama:
            messages.error(request, 'Nama kelas wajib diisi.')
        elif Kelas.objects.filter(nama__iexact=nama).exists():
            messages.error(request, 'Kelas tersebut sudah ada.')
        else:
            Kelas.objects.create(nama=nama)
            messages.success(request, 'Kelas berhasil ditambahkan.')
            return redirect('data_kelas')

    return render(request, 'students/tambah_kelas.html')

def edit_kelas(request, id):
    kelas = get_object_or_404(Kelas, id=id)

    if request.method == 'POST':
        nama = request.POST.get('nama', '').strip()

        if not nama:
            messages.error(request, 'Nama kelas wajib diisi.')

        elif Kelas.objects.filter(
            nama__iexact=nama
        ).exclude(id=kelas.id).exists():
            messages.error(request, 'Nama kelas tersebut sudah digunakan.')

        else:
            kelas.nama = nama
            kelas.save()

            messages.success(
                request,
                'Data kelas berhasil diperbarui.'
            )

            return redirect('data_kelas')

    return render(
        request,
        'students/edit_kelas.html',
        {
            'kelas': kelas
        }
    )

# =========================================================
# DATA GURU
# =========================================================

def data_guru(request):
    gurus = Guru.objects.all().order_by('nama')

    context = {
        'gurus': gurus,
    }

    return render(
        request,
        'students/data_guru.html',
        context
    )


def tambah_guru(request):

    if request.method == 'POST':

        nama = request.POST.get(
            'nama',
            ''
        ).strip()

        jenis_kelamin = request.POST.get(
            'jenis_kelamin',
            ''
        ).strip()

        if not nama:

            messages.error(
                request,
                'Nama guru wajib diisi.'
            )

        elif jenis_kelamin not in ['L', 'P']:

            messages.error(
                request,
                'Jenis kelamin wajib dipilih.'
            )

        elif Guru.objects.filter(
            nama__iexact=nama
        ).exists():

            messages.error(
                request,
                'Guru tersebut sudah ada.'
            )

        else:

            Guru.objects.create(
                nama=nama,
                jenis_kelamin=jenis_kelamin
            )

            messages.success(
                request,
                'Data guru berhasil ditambahkan.'
            )

            return redirect('data_guru')

    return render(
        request,
        'students/tambah_guru.html'
    )


def edit_guru(request, id):

    guru = get_object_or_404(
        Guru,
        id=id
    )

    if request.method == 'POST':

        nama = request.POST.get(
            'nama',
            ''
        ).strip()

        jenis_kelamin = request.POST.get(
            'jenis_kelamin',
            ''
        ).strip()

        if not nama:

            messages.error(
                request,
                'Nama guru wajib diisi.'
            )

        elif jenis_kelamin not in ['L', 'P']:

            messages.error(
                request,
                'Jenis kelamin wajib dipilih.'
            )

        elif Guru.objects.filter(
            nama__iexact=nama
        ).exclude(
            id=guru.id
        ).exists():

            messages.error(
                request,
                'Nama guru tersebut sudah digunakan.'
            )

        else:

            guru.nama = nama
            guru.jenis_kelamin = jenis_kelamin
            guru.save()

            messages.success(
                request,
                'Data guru berhasil diperbarui.'
            )

            return redirect('data_guru')

    return render(
        request,
        'students/edit_guru.html',
        {
            'guru': guru
        }
    )

def data_mata_pelajaran(request):
    mata_pelajarans = (
        MataPelajaran.objects
        .select_related('tahun_ajaran')
        .prefetch_related('kelas')
        .order_by('tahun_ajaran__nama', 'nama')
    )

    context = {
        'mata_pelajarans': mata_pelajarans,
    }

    return render(
        request,
        'students/data_mata_pelajaran.html',
        context
    )

def tambah_mata_pelajaran(request):
    tahun_ajarans = TahunAjaran.objects.all().order_by('-nama')
    kelass = Kelas.objects.all().order_by('nama')

    if request.method == 'POST':

        nama = request.POST.get('nama', '').strip()
        kitab = request.POST.get('kitab', '').strip()
        tahun_ajaran_id = request.POST.get('tahun_ajaran')
        kelas_ids = request.POST.getlist('kelas')
        aktif = request.POST.get('aktif') == 'on'

        # Validasi nama
        if not nama:
            messages.error(
                request,
                'Nama mata pelajaran wajib diisi.'
            )

        # Validasi tahun ajaran
        elif not tahun_ajaran_id:
            messages.error(
                request,
                'Tahun ajaran wajib dipilih.'
            )

        # Validasi kelas
        elif not kelas_ids:
            messages.error(
                request,
                'Minimal pilih satu kelas.'
            )

        else:

            try:
                tahun_ajaran = TahunAjaran.objects.get(
                    id=tahun_ajaran_id
                )

                # Cek mata pelajaran dengan nama yang sama
                # pada tahun ajaran yang sama
                if MataPelajaran.objects.filter(
                    nama__iexact=nama,
                    tahun_ajaran=tahun_ajaran
                ).exists():

                    messages.error(
                        request,
                        'Mata pelajaran tersebut sudah ada '
                        'pada tahun ajaran yang dipilih.'
                    )

                else:

                    mata_pelajaran = MataPelajaran.objects.create(
                        nama=nama,
                        kitab=kitab,
                        tahun_ajaran=tahun_ajaran,
                        aktif=aktif
                    )

                    # Simpan kelas yang berlaku
                    mata_pelajaran.kelas.set(kelas_ids)

                    messages.success(
                        request,
                        'Mata pelajaran berhasil ditambahkan.'
                    )

                    return redirect('data_mata_pelajaran')

            except TahunAjaran.DoesNotExist:

                messages.error(
                    request,
                    'Tahun ajaran tidak ditemukan.'
                )

    context = {
        'tahun_ajarans': tahun_ajarans,
        'kelass': kelass,
    }

    return render(
        request,
        'students/tambah_mata_pelajaran.html',
        context
    )

def edit_mata_pelajaran(request, id):
    mata_pelajaran = get_object_or_404(MataPelajaran, id=id)

    tahun_ajarans = TahunAjaran.objects.all().order_by('-nama')
    kelass = Kelas.objects.all().order_by('nama')

    if request.method == 'POST':

        nama = request.POST.get('nama', '').strip()
        kitab = request.POST.get('kitab', '').strip()
        tahun_ajaran_id = request.POST.get('tahun_ajaran')
        kelas_ids = request.POST.getlist('kelas')
        aktif = request.POST.get('aktif') == 'on'

        if not nama:
            messages.error(request, 'Nama mata pelajaran wajib diisi.')

        elif not tahun_ajaran_id:
            messages.error(request, 'Tahun ajaran wajib dipilih.')

        elif not kelas_ids:
            messages.error(request, 'Minimal pilih satu kelas.')

        else:
            try:
                tahun_ajaran = TahunAjaran.objects.get(
                    id=tahun_ajaran_id
                )

                sudah_ada = MataPelajaran.objects.filter(
                    nama__iexact=nama,
                    tahun_ajaran=tahun_ajaran
                ).exclude(
                    id=mata_pelajaran.id
                ).exists()

                if sudah_ada:
                    messages.error(
                        request,
                        'Mata pelajaran tersebut sudah ada '
                        'pada tahun ajaran yang dipilih.'
                    )

                else:
                    mata_pelajaran.nama = nama
                    mata_pelajaran.kitab = kitab
                    mata_pelajaran.tahun_ajaran = tahun_ajaran
                    mata_pelajaran.aktif = aktif
                    mata_pelajaran.save()

                    mata_pelajaran.kelas.set(kelas_ids)

                    messages.success(
                        request,
                        'Mata pelajaran berhasil diperbarui.'
                    )

                    return redirect('data_mata_pelajaran')

            except TahunAjaran.DoesNotExist:
                messages.error(
                    request,
                    'Tahun ajaran tidak ditemukan.'
                )

    context = {
        'mata_pelajaran': mata_pelajaran,
        'tahun_ajarans': tahun_ajarans,
        'kelass': kelass,
    }

    return render(
        request,
        'students/edit_mata_pelajaran.html',
        context
    )

# =========================================================
# RAPOT
# =========================================================

def rapot(request):

    # =====================================================
    # DATA MASTER
    # =====================================================

    tahun_ajarans = (
        TahunAjaran.objects
        .all()
        .order_by('-nama')
    )

    kelass = (
        Kelas.objects
        .all()
        .order_by('nama')
    )

    # =====================================================
    # FILTER UTAMA
    # =====================================================

    tahun_ajaran_id = request.GET.get(
        'tahun_ajaran',
        ''
    ).strip()

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    # =====================================================
    # PENCARIAN SISWA
    # =====================================================

    search = request.GET.get(
        'search',
        ''
    ).strip()

    # =====================================================
    # DATA SISWA
    # =====================================================

    students = Student.objects.none()

    if tahun_ajaran_id and kelas_id:

        students = (
            Student.objects
            .select_related(
                'kelas',
                'tahun_ajaran'
            )
            .filter(
                status=True,
                kelas_id=kelas_id,
                tahun_ajaran_id=tahun_ajaran_id,
            )
            .order_by('nama')
        )

        # -------------------------------------------------
        # CARI NAMA ATAU NIM
        # -------------------------------------------------

        if search:

            students = students.filter(
                Q(nama__icontains=search) |
                Q(nim__icontains=search)
            )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'tahun_ajarans':
            tahun_ajarans,

        'kelass':
            kelass,

        'students':
            students,

        'tahun_ajaran_id':
            tahun_ajaran_id,

        'kelas_id':
            kelas_id,

        'search':
            search,
    }

    return render(
        request,
        'students/rapot.html',
        context
    )

# =========================================================
# CETAK RAPOT PER SISWA
# =========================================================

def cetak_rapot_siswa(request, id):

    # =====================================================
    # AMBIL PARAMETER
    # =====================================================

    tahun_ajaran_id = request.GET.get(
        'tahun_ajaran',
        ''
    ).strip()

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    # =====================================================
    # AMBIL SISWA
    # =====================================================

    student = get_object_or_404(
        Student.objects.select_related(
            'kelas',
            'tahun_ajaran'
        ),
        id=id,
        status=True
    )

    # =====================================================
    # PASTIKAN SESUAI TAHUN AJARAN
    # DAN KELAS YANG DIPILIH
    # =====================================================

    if tahun_ajaran_id:

        if str(student.tahun_ajaran_id) != str(
            tahun_ajaran_id
        ):
            messages.error(
                request,
                'Siswa tidak sesuai dengan tahun ajaran yang dipilih.'
            )

            return redirect('rapot')

    if kelas_id:

        if str(student.kelas_id) != str(
            kelas_id
        ):
            messages.error(
                request,
                'Siswa tidak sesuai dengan kelas yang dipilih.'
            )

            return redirect('rapot')

    # =====================================================
    # TAHUN AJARAN
    # =====================================================

    tahun_ajaran = None

    if tahun_ajaran_id:

        tahun_ajaran = (
            TahunAjaran.objects
            .filter(
                id=tahun_ajaran_id
            )
            .first()
        )

    elif student.tahun_ajaran_id:

        tahun_ajaran = student.tahun_ajaran

    # =====================================================
    # KELAS
    # =====================================================

    kelas = student.kelas

    # =====================================================
    # NILAI
    # =====================================================

    penilaians = Penilaian.objects.none()

    if tahun_ajaran and kelas:

        penilaians = (
            Penilaian.objects
            .filter(
                student=student,
                mata_pelajaran__tahun_ajaran=tahun_ajaran,
                mata_pelajaran__kelas=kelas,
                mata_pelajaran__aktif=True,
            )
            .select_related(
                'student',
                'mata_pelajaran',
            )
            .order_by(
                'mata_pelajaran__nama'
            )
            .distinct()
        )

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {

        'student':
            student,

        'tahun_ajaran':
            tahun_ajaran,

        'kelas':
            kelas,

        'penilaians':
            penilaians,
    }

    return render(
        request,
        'students/cetak_rapot.html',
        context
    )

def import_siswa(request):
    if request.method != 'POST':
        return redirect('student_list')

    file = request.FILES.get('file')

    if not file:
        messages.error(request, 'Silakan pilih file Excel.')
        return redirect('student_list')

    if not file.name.lower().endswith('.xlsx'):
        messages.error(request, 'File harus berformat Excel .xlsx.')
        return redirect('student_list')

    try:
        workbook = openpyxl.load_workbook(file, data_only=True)
        sheet = workbook.active

        headers = [
            cell.value
            for cell in sheet[1]
        ]

        headers = [
            str(header).strip() if header is not None else ''
            for header in headers
        ]

        required_headers = [
            'Nama',
            'NIM',
            'JK',
            'Tempat Lahir',
            'Tanggal Lahir',
            'Alamat',
            'Semester',
            'Kelas',
            'Tahun Ajaran',
            'Prodi',
        ]

        for header in required_headers:
            if header not in headers:
                messages.error(
                    request,
                    f'Kolom "{header}" tidak ditemukan dalam file Excel.'
                )
                return redirect('student_list')

        header_index = {
            header: index
            for index, header in enumerate(headers)
        }

        berhasil = 0
        dilewati = 0

        for row in sheet.iter_rows(min_row=2, values_only=True):

            if not any(row):
                continue

            def get_value(nama_kolom):
                index = header_index.get(nama_kolom)

                if index is None or index >= len(row):
                    return ''

                value = row[index]

                if value is None:
                    return ''

                return str(value).strip()

            nama = get_value('Nama')
            nama_ayah = get_value('Nama Ayah')
            nim = get_value('NIM')
            if nim.endswith('.0'):
                nim = nim[:-2]
            jk = get_value('JK')
            tempat_lahir = get_value('Tempat Lahir')
            tanggal_lahir = row[
                header_index['Tanggal Lahir']
            ]
            alamat = get_value('Alamat')
            no_tlpn_wa = get_value('No. Telepon/WA')
            status = get_value('Status')
            asrama_nama = get_value('Asrama')
            semester = get_value('Semester')
            kelas_nama = get_value('Kelas')
            tahun_ajaran_nama = get_value('Tahun Ajaran')
            prodi = get_value('Prodi')

            # Lewati baris contoh dari template
            if nama.upper().startswith('CONTOH'):
                continue

            # Validasi data wajib
            if not nama or not nim:
                dilewati += 1
                continue

            # NIM harus unik
            if Student.objects.filter(nim=nim).exists():
                dilewati += 1
                continue

            # JK
            if jk not in ['L', 'P']:
                dilewati += 1
                continue

            # Tanggal lahir
            try:
                if isinstance(tanggal_lahir, datetime):
                    tanggal_lahir = tanggal_lahir.date()

                elif hasattr(tanggal_lahir, 'year'):
                    pass

                else:
                    tanggal_lahir = datetime.strptime(
                        str(tanggal_lahir),
                        '%d/%m/%Y'
                    ).date()

            except (ValueError, TypeError):
                dilewati += 1
                continue

            # Semester
            try:
                semester = int(float(semester))
            except (ValueError, TypeError):
                dilewati += 1
                continue

            # Kelas
            kelas = None

            if kelas_nama:
                kelas = Kelas.objects.filter(
                    nama__iexact=kelas_nama
                ).first()

                if not kelas:
                    dilewati += 1
                    continue

            # Tahun Ajaran
            tahun_ajaran = TahunAjaran.objects.filter(
                nama__iexact=tahun_ajaran_nama
            ).first()

            if not tahun_ajaran:
                dilewati += 1
                continue

            # Asrama
            asrama_master = None

            if asrama_nama:
                asrama_master = Asrama.objects.filter(
                    nama__iexact=asrama_nama
                ).first()

                if not asrama_master:
                    dilewati += 1
                    continue

            # Status
            status_value = status.lower() == 'aktif'

            Student.objects.create(
                nama=nama,
                nama_ayah=nama_ayah,
                nim=nim,
                jk=jk,
                tempat_lahir=tempat_lahir,
                tanggal_lahir=tanggal_lahir,
                alamat=alamat,
                no_tlpn_wa=no_tlpn_wa,
                status=status_value,
                asrama=asrama_nama,
                asrama_master=asrama_master,
                semester=semester,
                kelas=kelas,
                tahun_ajaran=tahun_ajaran,
                prodi=prodi,
            )

            berhasil += 1

        messages.success(
            request,
            f'Import selesai. {berhasil} siswa berhasil diimpor, '
            f'{dilewati} baris dilewati.'
        )

    except Exception as e:
        messages.error(
            request,
            f'Import gagal: {str(e)}'
        )

    return redirect('student_list')

def export_siswa(request):

    # =========================================
    # AMBIL FILTER DARI URL
    # =========================================

    search = request.GET.get(
        'search',
        ''
    ).strip()

    tahun_ajaran_id = request.GET.get(
        'tahun_ajaran',
        ''
    ).strip()

    kelas_id = request.GET.get(
        'kelas',
        ''
    ).strip()

    asrama_id = request.GET.get(
        'asrama',
        ''
    ).strip()

    status = request.GET.get(
        'status',
        ''
    ).strip()

    # =========================================
    # QUERY SISWA
    # =========================================

    students = (
        Student.objects
        .select_related(
            'kelas',
            'asrama_master',
            'tahun_ajaran'
        )
        .all()
        .order_by('nama')
    )

    # =========================================
    # SEARCH
    # =========================================

    if search:

        students = students.filter(
            Q(nama__icontains=search) |
            Q(nim__icontains=search)
        )

    # =========================================
    # FILTER TAHUN AJARAN
    # =========================================

    if tahun_ajaran_id:

        students = students.filter(
            tahun_ajaran_id=tahun_ajaran_id
        )

    # =========================================
    # FILTER KELAS
    # =========================================

    if kelas_id:

        students = students.filter(
            kelas_id=kelas_id
        )

    # =========================================
    # FILTER ASRAMA
    # =========================================

    if asrama_id:

        students = students.filter(
            asrama_master_id=asrama_id
        )

    # =========================================
    # FILTER STATUS
    # =========================================

    if status == 'aktif':

        students = students.filter(
            status=True
        )

    elif status == 'nonaktif':

        students = students.filter(
            status=False
        )

    # =========================================
    # BUAT EXCEL
    # =========================================

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "Data Siswa"

    # =========================================
    # HEADER
    # =========================================

    headers = [
        'Nama',
        'Nama Ayah',
        'NIM',
        'JK',
        'Tempat Lahir',
        'Tanggal Lahir',
        'Alamat',
        'No. Telepon/WA',
        'Status',
        'Asrama',
        'Semester',
        'Kelas',
        'Tahun Ajaran',
        'Prodi',
    ]

    sheet.append(headers)

    # =========================================
    # DATA
    # =========================================

    for student in students:

        sheet.append([

            student.nama,

            student.nama_ayah,

            student.nim,

            student.jk,

            student.tempat_lahir,

            student.tanggal_lahir,

            student.alamat,

            student.no_tlpn_wa,

            (
                'Aktif'
                if student.status
                else 'Nonaktif'
            ),

            (
                student.asrama_master.nama
                if student.asrama_master
                else student.asrama
            ),

            student.semester,

            (
                student.kelas.nama
                if student.kelas
                else ''
            ),

            (
                student.tahun_ajaran.nama
                if student.tahun_ajaran
                else ''
            ),

            student.prodi,

        ])

    # =========================================
    # RESPONSE
    # =========================================

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; '
        'filename="data_siswa_filter.xlsx"'
    )

    workbook.save(response)

    return response

def download_format_siswa(request):

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data Siswa"

    headers = [
        'Nama',
        'Nama Ayah',
        'NIM',
        'JK',
        'Tempat Lahir',
        'Tanggal Lahir',
        'Alamat',
        'No. Telepon/WA',
        'Status',
        'Asrama',
        'Semester',
        'Kelas',
        'Tahun Ajaran',
        'Prodi',
    ]

    sheet.append(headers)

    # Contoh format
    sheet.append([
        'CONTOH - HAPUS BARIS INI',
        'Ahmad',
        '10001',
        'L',
        'Kediri',
        '01/01/2005',
        'Alamat siswa',
        '081234567890',
        'Aktif',
        'Asrama A',
        1,
        '1A',
        '2026',
        'Manajemen',
    ])

    # Lebar kolom
    widths = {
        'A': 30,
        'B': 25,
        'C': 15,
        'D': 10,
        'E': 20,
        'F': 18,
        'G': 40,
        'H': 20,
        'I': 14,
        'J': 20,
        'K': 12,
        'L': 14,
        'M': 18,
        'N': 30,
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    # Freeze header
    sheet.freeze_panes = 'A2'

    # Filter
    sheet.auto_filter.ref = f"A1:N{sheet.max_row}"

    # Style header
    from openpyxl.styles import Font, PatternFill, Alignment

    for cell in sheet[1]:
        cell.font = Font(
            bold=True,
            color='FFFFFF'
        )

        cell.fill = PatternFill(
            fill_type='solid',
            fgColor='1F2937'
        )

        cell.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

    # Response download
    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; filename="format_data_siswa.xlsx"'
    )

    workbook.save(response)

    return response